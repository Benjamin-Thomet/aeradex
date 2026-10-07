import io
from decimal import Decimal

import pytest

from aeradex import api, berichte, check, reports
from aeradex.book import Book, BookError
from aeradex.testing import make_book

D = Decimal


@pytest.fixture
def book(tmp_path):
    b = make_book(tmp_path, eroeffnung={"1020": 20000, "2800": -20000})
    for m in (1, 2, 3):
        api.post_entry(Book(b.root), f"2026-{m:02d}-05", "1020", "3400", 5000 + m * 100, "Umsatz")
        api.post_entry(Book(b.root), f"2026-{m:02d}-25", "6000", "1020", "1200", "Miete")
    return Book(b.root)


def test_exports(book):
    data, name = api.report_export(book, "erfolgsrechnung", "pdf", jahr=2026, spalten="monat", periode="q1")
    assert data.startswith(b"%PDF") and name.startswith("Erfolgsrechnung")
    pytest.importorskip("openpyxl")
    from openpyxl import load_workbook
    data, name = api.report_export(book, "erfolgsrechnung", "xlsx", jahr=2026, periode="q1", vergleich="vorperiode")
    ws = load_workbook(io.BytesIO(data)).active
    values = {ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(1, ws.max_row + 1)}
    assert values["Betriebsertrag aus Lieferungen und Leistungen"] == D("15600")    # numbers, not text
    assert any("Stand" in str(ws.cell(r, 1).value) for r in range(1, ws.max_row + 1))
    data, _ = api.report_export(book, "bilanz", "csv", jahr=2026)
    assert "Total Aktiven" in data.decode("utf-8-sig")
    with pytest.raises(BookError, match="Format"):
        api.report_export(book, "bilanz", "docx", jahr=2026)


def test_text_output(book):
    out = berichte.text(reports.run(book, "erfolgsrechnung", jahr=2026, periode="q1"), detail=True)
    assert "Betriebsertrag" in out and "3400 Dienstleistungserlöse" in out and "15'600.00" in out


def test_templates(book):
    api.report_template_save(book, "Quartal Treuhand", "erfolgsrechnung", periode="q1", spalten="monat",
                             vergleich="vorjahr", jahr=2026)
    t = berichte.templates(Book(book.root))["Quartal Treuhand"]
    assert t == {"typ": "erfolgsrechnung", "parameter": {"periode": "q1", "spalten": "monat", "vergleich": "vorjahr"}}
    rep = api.report(Book(book.root), vorlage="Quartal Treuhand", jahr=2026)
    assert rep["vorlage"] == "Quartal Treuhand" and len(rep["spalten"]) == 7       # 3 months, total, ref, diff, pct
    api.report_template_delete(Book(book.root), "Quartal Treuhand")
    assert berichte.templates(Book(book.root)) == {}


def test_comment_turns_stale(book):
    api.report_comment_save(book, "Umsatz steigt jeden Monat.", "erfolgsrechnung", autor="Agent (Test)",
                            jahr=2026, periode="q1")
    rep = api.report(Book(book.root), "erfolgsrechnung", jahr=2026, periode="q1")
    assert rep["kommentar"]["text"] == "Umsatz steigt jeden Monat." and rep["kommentar"]["veraltet"] is False
    assert api.report(Book(book.root), "erfolgsrechnung", jahr=2026, periode="q2")["kommentar"] is None
    api.post_entry(Book(book.root), "2026-03-30", "6500", "1020", "50", "Papier")
    assert api.report(Book(book.root), "erfolgsrechnung", jahr=2026, periode="q1")["kommentar"]["veraltet"] is True


def test_monthly_package_and_mail(book, monkeypatch):
    sent = {}

    class FakeSMTP:
        def __init__(self, host, port, timeout=None):
            sent["host"] = host

        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

        def starttls(self, context=None):
            sent["tls"] = True

        def login(self, user, pw):
            sent["login"] = user

        def send_message(self, msg):
            sent["to"] = msg["To"]
            sent["files"] = [p.get_filename() for p in msg.iter_attachments()]

    import smtplib
    monkeypatch.setattr(smtplib, "SMTP", FakeSMTP)
    monkeypatch.setenv("AERADEX_SMTP_PASSWORD", "geheim")
    with pytest.raises(BookError, match="nicht eingerichtet"):
        api.monthly_report(book, "2026-03", mail=True)
    b = Book(book.root)
    b.settings.data["mail"] = {"smtp": "mail.example.ch", "von": "bh@example.ch", "an": ["chefin@example.ch"]}
    b.save_settings()
    out = api.monthly_report(Book(book.root), "2026-03", mail=True)
    assert out["gesendet_an"] == "chefin@example.ch" and sent["tls"] and sent["login"] == "bh@example.ch"
    assert "Monatsbericht 2026-03.pdf" in sent["files"]
    assert (book.root / "berichte/2026-03/Monatsbericht 2026-03.pdf").exists()


def test_budget_errors_block_the_book(book):
    p = book.root / "budget" / "2026.yaml"
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text('konten:\n  "1020": {jahr: 10}\n')
    assert any(i.level == "fehler" and "Erfolgskonto" in i.message for i in check.run(Book(book.root)))


def test_report_list_and_unknown(book):
    names = [r["name"] for r in api.report_list(book)["berichte"]]
    assert {"erfolgsrechnung", "bilanz", "geldfluss", "kennzahlen", "debitoren", "kreditoren", "umsatz"} <= set(names)


def test_with_git(tmp_path):
    b = make_book(tmp_path, eroeffnung={"1020": 1000, "2800": -1000}, git=True)
    api.post_entry(Book(b.root), "2026-01-05", "1020", "3400", "500", "Umsatz")
    assert api.report_template_save(Book(b.root), "Monat", "erfolgsrechnung", spalten="monat")["commit"]
    assert api.report_comment_save(Book(b.root), "Gut.", "erfolgsrechnung", jahr=2026)["commit"]
    assert api.budget_set(Book(b.root), 2026, "3400", 6000)["commit"]
    out = api.monthly_report(Book(b.root), "2026-01")
    assert out["ok"] and (b.root / "berichte/2026-01/Monatsbericht 2026-01.pdf").exists()
    import subprocess
    status = subprocess.run(["git", "-C", str(b.root), "status", "--porcelain"], capture_output=True, text=True).stdout
    assert "berichte" not in status                                   # the package is output, not part of the book
    assert berichte.stand(Book(b.root))["commit"]


def test_dossier_has_cash_flow_and_kpis(book):
    import zipfile
    from pathlib import Path
    res = api.dossier(book, 2026, ["geldfluss"], ["pdf", "csv"])
    names = zipfile.ZipFile(Path(res["datei"])).namelist()
    assert any("Geldfluss und Kennzahlen 2026.pdf" in n for n in names)
    assert any("Kennzahlen 2026.csv" in n for n in names)


def test_agent_comment_with_fake_backend(book, monkeypatch):
    pytest.importorskip("starlette")
    from aeradex.web import chat
    prompts = []

    class FakeAgent:
        def __init__(self, root):
            self.root = root

        def run(self, events, prompt):
            prompts.append(prompt)
            api.report_comment_save(Book(self.root), "Umsatz Q1 15'600, Miete stabil.", "erfolgsrechnung",
                                    autor="Agent (Test)", jahr=2026, periode="q1")
            events.put({"type": "text", "text": "Kommentar gespeichert."})

    monkeypatch.setattr(chat, "backend", lambda preferred=None: "fake")
    monkeypatch.setitem(chat.AGENTS, "fake", FakeAgent)
    said = berichte.run_agent_comment(book.root, "erfolgsrechnung", jahr=2026, periode="q1")
    assert said == "Kommentar gespeichert."
    assert 'period_report(typ="erfolgsrechnung", jahr=2026, periode="q1")' in prompts[0]
    assert 'save_report_comment(typ="erfolgsrechnung", jahr=2026, periode="q1", text=…)' in prompts[0]
    assert "rechne nichts selbst" in prompts[0]
    rep = api.report(Book(book.root), "erfolgsrechnung", jahr=2026, periode="q1")
    assert rep["kommentar"]["autor"] == "Agent (Test)"


def test_template_and_plain_report_share_the_comment(book):
    api.report_template_save(book, "Q1", "erfolgsrechnung", periode="q1")
    api.report_comment_save(Book(book.root), "Gilt für beide.", vorlage="Q1", jahr=2026)
    assert api.report(Book(book.root), "erfolgsrechnung", jahr=2026, periode="q1")["kommentar"]["text"] == "Gilt für beide."
