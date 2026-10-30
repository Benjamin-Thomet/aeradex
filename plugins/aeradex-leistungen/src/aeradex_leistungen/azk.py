"""The SMGV Arbeitszeitkontrolle (GAV Maler- und Gipsergewerbe, Art. 8.9): the association's Excel file, filled in per
person and year from the recorded hours — work per day, travel time, absences, holidays, Pensum and the employee's data.

The file stays the association's: its formulas compute the balances, surcharges and holiday account. aeradex only
writes the yellow input cells, straight into the sheet XML of the workbook (protection, formatting, validation and the
hidden helper sheets stay exactly as delivered; a library that loads and saves the workbook would drop parts of it)
and asks Excel to recalculate everything when the file is opened.

The template lies in the book (`leistungen/azk/vorlage-<JJJJ>.xlsx`, uploaded under Leistungen › Arbeitszeit); it is
not shipped with aeradex. A year the association has not published yet is made from the last one: only the year cell
is a constant, the calendar follows from it. Which labels and codes are where is checked against the template, so a
changed layout of a new version is reported instead of filling the wrong cells.
"""
from __future__ import annotations

import io
import re
import zipfile
from collections import defaultdict
from datetime import date, timedelta
from decimal import Decimal
from pathlib import Path
from xml.sax.saxutils import escape

from aeradex import payroll
from aeradex.book import Book, BookError

from . import abwesenheit, daten

ZERO = Decimal("0")
MONTHS = ["JAN", "FEB", "MRZ", "APR", "MAI", "JUN", "JUL", "AUG", "SEP", "OKT", "NOV", "DEZ"]
FIRST_DAY_COL = 6                 # F = day 1 … AJ = day 31
ROW_DAY = 24                      # 06.00–20.00 h
ROW_NIGHT_EARLY, ROW_NIGHT_LATE = 23, 25
ROW_TRAVEL = 26                   # Reisezeit (gesamt)
CODE_COL = 41                     # AO: the absence code of the rows below
EXPECTED = {("H_Label", "C257"): "00.00-06.00h", ("H_Label", "C258"): "06.00-20.00h",
            ("H_Label", "C259"): "20.00-24.00h", ("H_Label", "C313"): "Reisezeit (gesamt)",
            ("H_Label", "C276"): "Ferien", ("H_Label", "C101"): "Mitarbeiter:in"}
# aeradex absence kinds → AZK codes (Wegleitung, «Übersicht Codes»)
ABSENCE_CODES = {"Ferien": "f", "Krank": "k", "Unfall": "u", "Militär/ZS": "m", "Mutterschaft/Vaterschaft": "e",
                 "Weiterbildung": "ku", "Kompensation": "kj", "Unbezahlt": "fu"}
CATEGORY_CODES = {"berufsschule": "bs", "schule": "bs", "überbetrieblicher kurs": "ku", "kurs": "ku",
                  "weiterbildung": "ku", "ferien": "f", "krank": "k"}
TRAVEL = "Reisezeit"


def folder(book: Book) -> Path:
    return book.root / "leistungen" / "azk"


def templates(book: Book) -> dict[int, Path]:
    out = {}
    for p in sorted(folder(book).glob("vorlage-*.xlsx")) if folder(book).exists() else []:
        if m := re.fullmatch(r"vorlage-(\d{4})", p.stem):
            out[int(m.group(1))] = p
    return out


# ---------- the workbook as XML ----------

def _col(n: int) -> str:
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def _col_index(ref: str) -> int:
    n = 0
    for ch in re.match(r"[A-Z]+", ref).group(0):
        n = n * 26 + ord(ch) - 64
    return n


def _sheet_paths(z: zipfile.ZipFile) -> dict[str, str]:
    wb = z.read("xl/workbook.xml").decode("utf-8")
    rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    targets = dict(re.findall(r'<Relationship[^>]*Id="([^"]+)"[^>]*Target="([^"]+)"', rels))
    targets.update({a: b for b, a in re.findall(r'<Relationship[^>]*Target="([^"]+)"[^>]*Id="([^"]+)"', rels)})
    out = {}
    for name, rid in re.findall(r'<sheet [^>]*name="([^"]+)"[^>]*r:id="([^"]+)"', wb):
        t = targets[rid]
        out[name.replace("&amp;", "&")] = t.lstrip("/") if t.startswith("/") else f"xl/{t}"
    return out


def _cell_xml(ref: str, style: str, value) -> str:
    s = f' s="{style}"' if style else ""
    if value is None or value == "":
        return f'<c r="{ref}"{s}/>'
    if isinstance(value, (int, float, Decimal)):
        return f'<c r="{ref}"{s}><v>{value}</v></c>'
    return f'<c r="{ref}"{s} t="inlineStr"><is><t xml:space="preserve">{escape(str(value))}</t></is></c>'


_CELL = r'<c r="{ref}"(?=[ />])[^>]*?(?:/>|>.*?</c>)'


def set_cells(xml: str, values: dict[str, object]) -> str:
    """Write values into the cells of a sheet's XML, keeping each cell's style; cells or rows that are not there yet
    are added in order. A formula in an input cell (the template's example text) is replaced by the value."""
    rows: dict[int, dict[str, object]] = defaultdict(dict)
    for ref, v in values.items():
        rows[int(re.search(r"\d+", ref).group(0))][ref] = v
    for r, cells in sorted(rows.items()):
        m = re.search(rf'<row r="{r}"(?=[ >/])[^>]*?(/>|>(.*?)</row>)', xml, re.S)
        if m is None:
            new_row = f'<row r="{r}">' + "".join(_cell_xml(ref, "", v) for ref, v in
                                                 sorted(cells.items(), key=lambda kv: _col_index(kv[0]))) + "</row>"
            later = [x for x in re.finditer(r'<row r="(\d+)"', xml) if int(x.group(1)) > r]
            if later:
                xml = xml[:later[0].start()] + new_row + xml[later[0].start():]
            else:
                xml = xml.replace("</sheetData>", new_row + "</sheetData>").replace("<sheetData/>",
                                                                                    f"<sheetData>{new_row}</sheetData>")
            continue
        if m.group(1) == "/>":
            open_tag = m.group(0)[:-2] + ">"
            body = ""
        else:
            open_tag = m.group(0)[:m.group(0).index(">") + 1]
            body = m.group(2)
        for ref, v in cells.items():
            cm = re.search(_CELL.format(ref=ref), body, re.S)
            if cm:
                old = cm.group(0)
                if re.search(r'<f[^>]*t="shared"[^>]*ref=', old):
                    raise BookError(f"Zelle {ref} der Vorlage ist keine Eingabezelle")
                style = (re.search(r' s="(\d+)"', old.split(">")[0]) or [None, ""])[1]
                body = body[:cm.start()] + _cell_xml(ref, style, v) + body[cm.end():]
            else:
                idx = _col_index(ref)
                after = [x for x in re.finditer(r'<c r="([A-Z]+)\d+"', body) if _col_index(x.group(1)) > idx]
                pos = after[0].start() if after else len(body)
                row_style = re.search(r' s="(\d+)"', open_tag) if 'customFormat="1"' in open_tag else None
                body = body[:pos] + _cell_xml(ref, row_style.group(1) if row_style else "", v) + body[pos:]
        xml = xml[:m.start()] + open_tag + body + "</row>" + xml[m.end():]
    return xml


def _recalculate_on_open(files: dict[str, bytes]) -> None:
    """Excel and LibreOffice compute every formula anew when the file is opened (the stored results are old)."""
    wb = files["xl/workbook.xml"].decode("utf-8")
    if "<calcPr" in wb:
        wb = re.sub(r"<calcPr([^>]*?)/>", lambda m: "<calcPr" + re.sub(r'\s*fullCalcOnLoad="[^"]*"', "", m.group(1))
                    + ' fullCalcOnLoad="1"/>', wb, count=1)
    else:                                      # CT_Workbook order: calcPr right after definedNames/sheets
        tag = '<calcPr fullCalcOnLoad="1"/>'
        for anchor in ("</definedNames>", "</externalReferences>", "</functionGroups>", "</sheets>"):
            if anchor in wb:
                wb = wb.replace(anchor, anchor + tag, 1)
                break
    files["xl/workbook.xml"] = wb.encode("utf-8")
    if "xl/calcChain.xml" in files:
        del files["xl/calcChain.xml"]
        rels = files["xl/_rels/workbook.xml.rels"].decode("utf-8")
        files["xl/_rels/workbook.xml.rels"] = re.sub(r'<Relationship[^>]*Target="calcChain.xml"[^>]*/>', "",
                                                     rels).encode("utf-8")
        ct = files["[Content_Types].xml"].decode("utf-8")
        files["[Content_Types].xml"] = re.sub(r'<Override[^>]*PartName="/xl/calcChain.xml"[^>]*/>', "", ct).encode("utf-8")


def write(data: bytes, sheets: dict[str, dict[str, object]]) -> bytes:
    """The workbook with values written into the named sheets' cells."""
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        files = {i.filename: z.read(i.filename) for i in z.infolist()}
        paths = _sheet_paths(z)
    for sheet, values in sheets.items():
        if sheet not in paths:
            raise BookError(f"Die Vorlage hat kein Blatt «{sheet}»")
        files[paths[sheet]] = set_cells(files[paths[sheet]].decode("utf-8"), values).encode("utf-8")
    _recalculate_on_open(files)
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for name in ["[Content_Types].xml"] + [n for n in files if n != "[Content_Types].xml"]:
            z.writestr(name, files[name])
    return out.getvalue()


# ---------- the template ----------

def inspect(data: bytes) -> dict:
    """Year, variant and the absence-code rows of a template; raises when it is not the expected layout."""
    import warnings
    import openpyxl
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            wb = openpyxl.load_workbook(io.BytesIO(data))            # read only: never saved with openpyxl
    except Exception:
        raise BookError("Das ist keine Excel-Datei der Arbeitszeitkontrolle") from None
    missing = [m for m in ["Admin_Konfig", "Basisblatt", "H_Label", *MONTHS] if m not in wb.sheetnames]
    if missing:
        raise BookError(f"Keine SMGV-Arbeitszeitkontrolle (es fehlt {', '.join(missing)})")
    for (sheet, ref), text in EXPECTED.items():
        if str(wb[sheet][ref].value or "").strip() != text:
            raise BookError(f"Die Vorlage hat ein anderes Format ({sheet}!{ref} ist nicht «{text}») — "
                            "diese Version kann aeradex noch nicht ausfüllen")
    codes = {}
    for r in range(30, 60):
        code = wb["JAN"].cell(r, CODE_COL).value
        if isinstance(code, str) and code.strip():
            codes[code.strip()] = r
    year = wb["Admin_Konfig"]["F3"].value
    if not isinstance(year, (int, float)) or not 2000 <= int(year) <= 2100:
        raise BookError("Die Vorlage nennt kein Jahr (Admin_Konfig F3)")
    return {"jahr": int(year), "codes": codes,
            "reisezeit": str(wb["Admin_Konfig"]["G17"].value or ""),
            "variante": "ohne Abzug Reisezeit" if wb["Admin_Konfig"]["G17"].value == "komplett vergütet"
            else "mit Abzug Reisezeit"}


def for_year(data: bytes, jahr: int) -> bytes:
    """The template of another year: only the year cell changes, the calendar follows from it."""
    return write(data, {"Admin_Konfig": {"F3": int(jahr)}})


def store_template(book: Book, data: bytes, jahr: int | None = None) -> tuple[dict, list[Path]]:
    """An uploaded template (for its own year, or made into `jahr`)."""
    info = inspect(data)
    if jahr and int(jahr) != info["jahr"]:
        data = for_year(data, int(jahr))
        info = {**info, "jahr": int(jahr), "abgeleitet": True}
    folder(book).mkdir(parents=True, exist_ok=True)
    target = folder(book) / f"vorlage-{info['jahr']}.xlsx"
    target.write_bytes(data)
    return {k: v for k, v in info.items() if k != "codes"} | {"datei": book.rel(target)}, [target]


def template(book: Book, jahr: int) -> bytes:
    """The template of the year; one of an earlier year is made into this year (the association may not have
    published it yet)."""
    have = templates(book)
    if jahr in have:
        return have[jahr].read_bytes()
    earlier = [y for y in have if y < jahr]
    if not earlier:
        raise BookError("Zuerst die Arbeitszeitkontrolle des SMGV (Excel) als Vorlage hochladen")
    return for_year(have[max(earlier)].read_bytes(), jahr)


# ---------- the person's year ----------

def _serial(d) -> int:
    return (d - date(1899, 12, 30)).days


def _daily_target(person: dict, d: date) -> Decimal:
    week = daten.week_hours(person, d) if person.get("wochen") else person.get("woche")
    return (Decimal(str(week)) / 5).quantize(Decimal("0.01")) if week else Decimal("8")


def year_days(book: Book, nr: str, jahr: int) -> dict[date, dict]:
    """Per day: work (06–20 h), travel, absence hours per AZK code — what goes into the month sheets."""
    p = daten.person(book, nr)
    out: dict[date, dict] = defaultdict(lambda: {"arbeit": ZERO, "reise": ZERO, "codes": defaultdict(lambda: ZERO)})
    for e in daten.entries(book):
        if e["wer"] != p["nummer"] or e["art"] != "Zeit" or e["datum"].year != jahr:
            continue
        cat = (e.get("kategorie") or "").strip()
        if cat == TRAVEL:
            out[e["datum"]]["reise"] += e["menge"]
        elif cat.lower() in CATEGORY_CODES and not e["abrechenbar"]:
            out[e["datum"]]["codes"][CATEGORY_CODES[cat.lower()]] += e["menge"]
        elif cat.lower() != "feiertag":
            out[e["datum"]]["arbeit"] += e["menge"]
    hols = abwesenheit.holidays(book, jahr)
    for a in abwesenheit.items(book):
        if a["wer"] != p["nummer"]:
            continue
        d = max(a["von"], date(jahr, 1, 1))
        while d <= min(a["bis"], date(jahr, 12, 31)):
            if d.weekday() < 5 and d not in hols:
                code = ABSENCE_CODES.get(a["art"], "a")
                out[d]["codes"][code] += (_daily_target(p, d) * Decimal(str(a["anteil"]))).quantize(Decimal("0.01"))
            d += timedelta(days=1)
    return dict(out)


def _pensum(book: Book, nr: str, emp: dict | None, d: date) -> Decimal:
    p = daten.person(book, nr)
    if p.get("pensum_ab"):
        ab = sorted((str(k), v) for k, v in p["pensum_ab"].items())
        current = [v for k, v in ab if k <= d.isoformat()]
        if current and p.get("lohn"):
            return Decimal(str(current[-1]))
    return Decimal(str((emp or {}).get("pensum") or 100))


def fill(book: Book, nr: str, jahr: int) -> bytes:
    """The Arbeitszeitkontrolle of a person and year, filled in."""
    data = template(book, jahr)
    info = inspect(data)
    p = daten.person(book, nr)
    emp = payroll.employees(book).get(p["nummer"]) if p.get("lohn") else None
    s = book.settings
    base: dict[str, object] = {"G3": s.firma, "G6": p["name"]}
    if emp:
        a = emp.get("adresse") or {}
        base.update({"Q6": " ".join(x for x in (a.get("strasse"), str(a.get("nr") or "")) if x),
                     "Q7": " ".join(x for x in (str(a.get("plz") or ""), a.get("ort")) if x),
                     "AA6": emp.get("telefon") or "", "AA7": emp.get("email") or "",
                     "G9": int(ahv) if (ahv := re.sub(r"\D", "", str(emp.get("ahv_nr") or ""))) else ""})
        for ref, key in (("Q9", "geburtsdatum"), ("AA9", "eintritt"), ("AE9", "austritt")):
            if emp.get(key):
                from aeradex.files import parse_date
                base[ref] = _serial(parse_date(emp[key]))
    else:
        base.update({ref: "" for ref in ("Q6", "Q7", "AA6", "AA7", "G9")})
    vortrag = (p.get("vortrag") or {}).get(jahr) or (p.get("vortrag") or {}).get(str(jahr))
    if vortrag not in (None, ""):
        base["K12"] = Decimal(str(vortrag))
    ferien = (p.get("ferien_vortrag") or {}).get(jahr) or (p.get("ferien_vortrag") or {}).get(str(jahr))
    if ferien not in (None, ""):
        base["K13"] = Decimal(str(ferien))
    for d in abwesenheit.holidays(book, jahr):
        base[f"{_col(2 + d.day)}{20 + d.month}"] = "x"
    for m in range(12):
        pct = _pensum(book, p["nummer"], emp, date(jahr, m + 1, 1))
        for col in ("E", "G", "I", "K", "M"):
            base[f"{col}{43 + m}"] = pct
    sheets: dict[str, dict[str, object]] = {"Basisblatt": base}
    days = year_days(book, p["nummer"], jahr)
    for d, v in days.items():
        col = _col(FIRST_DAY_COL + d.day - 1)
        cells = sheets.setdefault(MONTHS[d.month - 1], {})
        if v["arbeit"]:
            cells[f"{col}{ROW_DAY}"] = v["arbeit"]
        if v["reise"]:
            cells[f"{col}{ROW_TRAVEL}"] = v["reise"]
        for code, hours in v["codes"].items():
            row = info["codes"].get(code) or info["codes"].get("a")
            if row and hours:
                cells[f"{col}{row}"] = hours
    return write(data, sheets)


def summary(book: Book, nr: str, jahr: int) -> list[dict]:
    """Per month: work, travel, absences in hours — what the Arbeitszeitkontrolle receives, to see in aeradex."""
    out = [{"monat": m + 1, "arbeit": ZERO, "reise": ZERO, "absenzen": ZERO, "tage": 0} for m in range(12)]
    for d, v in year_days(book, nr, jahr).items():
        row = out[d.month - 1]
        row["arbeit"] += v["arbeit"]
        row["reise"] += v["reise"]
        row["absenzen"] += sum(v["codes"].values(), ZERO)
        row["tage"] += 1 if v["arbeit"] or v["reise"] else 0
    return out
