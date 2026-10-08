"""Bookings from a spreadsheet: an Excel template to fill in, and a reader that turns a filled-in
.xlsx or .csv into rows for the booking grid. Nothing is booked here — the rows land in the grid,
where the person checks them and books them all at once (`api.post_entries`)."""
from __future__ import annotations

import csv
import io
from datetime import date, datetime
from decimal import Decimal

from .book import Book, BookError

# Grid columns in order; each with the headings accepted in a file (lower case, without dots).
COLUMNS = [("datum", ("datum", "date", "buchungsdatum", "belegdatum")),
           ("text", ("beschreibung", "text", "buchungstext", "bezeichnung")),
           ("soll", ("soll", "sollkonto", "konto soll", "debit")),
           ("haben", ("haben", "habenkonto", "konto haben", "credit")),
           ("betrag", ("betrag", "betrag chf", "amount", "chf")),
           ("mwst", ("mwst", "mwst-code", "mwst code", "steuercode", "vat"))]
MAX_ROWS = 2000


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, float):
        value = Decimal(str(value))
    if isinstance(value, (int, Decimal)):
        return f"{Decimal(value):f}"
    return str(value).strip()


def _header_map(row: list[str]) -> dict[str, int] | None:
    names = [str(c or "").strip().lower().replace(".", "") for c in row]
    found = {}
    for key, aliases in COLUMNS:
        for i, n in enumerate(names):
            if n in aliases:
                found[key] = i
                break
    return found if {"soll", "haben", "betrag"} <= found.keys() else None


def _rows_from_table(table: list[list]) -> list[list[str]]:
    header, start = None, 0
    for i, row in enumerate(table[:20]):              # a title or a note may sit above the heading
        header = _header_map([_cell(c) for c in row])
        if header:
            start = i + 1
            break
    if header is None:
        raise BookError("Keine Kopfzeile gefunden: erwartet werden mindestens die Spalten Soll, Haben und Betrag "
                        "(am einfachsten die Vorlage verwenden)")
    out = []
    for row in table[start:]:
        cells = [_cell(row[header[k]]) if k in header and header[k] < len(row) else "" for k, _ in COLUMNS]
        if not any(cells[1:5]):
            continue
        # accounts from the template's dropdown come as '1020 Bank' — the grid wants the number first anyway
        out.append(cells)
        if len(out) > MAX_ROWS:
            raise BookError(f"Mehr als {MAX_ROWS} Zeilen — bitte die Datei aufteilen")
    if not out:
        raise BookError("Die Datei enthält keine Buchungszeilen")
    return out


def read(data: bytes, filename: str) -> list[list[str]]:
    """Rows [datum, text, soll, haben, betrag, mwst] as text, in the order of the grid."""
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        try:
            import openpyxl
        except ImportError:
            raise BookError("Für Excel-Dateien fehlt openpyxl: pip install 'aeradex[excel]' — oder als CSV speichern") from None
        try:
            sheet = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
        except Exception as exc:
            raise BookError(f"Die Excel-Datei lässt sich nicht lesen: {exc}") from None
        return _rows_from_table([list(r) for r in sheet.iter_rows(values_only=True)])
    if name.endswith((".csv", ".txt")):
        text = data.decode("utf-8-sig", errors="replace")
        sample = text[:4096]
        # Swiss Excel writes ';' — count rather than sniff, a title line above the heading confuses the sniffer
        delimiter = max(";\t,", key=sample.count)
        return _rows_from_table(list(csv.reader(io.StringIO(text), delimiter=delimiter)))
    if name.endswith(".xls"):
        raise BookError("Altes Excel-Format (.xls): bitte in Excel als .xlsx speichern")
    raise BookError("Erwartet wird eine Excel-Datei (.xlsx) oder CSV")


def template(book: Book) -> bytes:
    """An .xlsx to fill in: the booking sheet with account and MWST dropdowns, the chart of accounts,
    and a short guide."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    from . import mwst

    vat = mwst.config(book)["methode"] != "keine"
    wb = Workbook()
    ws = wb.active
    ws.title = "Buchungen"
    heads = ["Datum", "Beschreibung", "Soll", "Haben", "Betrag"] + (["MWST"] if vat else [])
    widths = [12, 44, 26, 26, 14, 10]
    ws.append(heads)
    bold, fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1C1917")
    for i, cell in enumerate(ws[1]):
        cell.font, cell.fill = bold, fill
        cell.alignment = Alignment(horizontal="right" if heads[i] == "Betrag" else "left")
        ws.column_dimensions[cell.column_letter].width = widths[i]
    ws.freeze_panes = "A2"

    accounts = sorted(book.accounts.values(), key=lambda a: a.nr)
    bank = book.settings.konto("bank")
    name = lambda nr: f"{nr} {book.accounts[nr].name}" if nr in book.accounts else nr  # noqa: E731
    first = next((a.nr for a in accounts if a.nr.startswith("6")), accounts[0].nr if accounts else "")
    today = date.today()
    examples = [[date(today.year, today.month, 1), "Beispiel: Büromaterial (Zeile löschen)", name(first), name(bank), 45.80]]
    for row in examples:
        ws.append(row + ([""] if vat else []))
    grey = Font(italic=True, color="78716C")
    for cell in ws[2]:
        cell.font = grey
    for r in range(2, 502):
        ws.cell(r, 1).number_format = "DD.MM.YYYY"
        ws.cell(r, 5).number_format = "#,##0.00"

    kp = wb.create_sheet("Konten")
    kp.append(["Konto", "Bezeichnung", "Klasse"])
    for c in kp[1]:
        c.font = Font(bold=True)
    for a in accounts:
        kp.append([f"{a.nr} {a.name}", a.name, a.klasse])
    kp.column_dimensions["A"].width, kp.column_dimensions["B"].width = 40, 36
    dv = DataValidation(type="list", formula1=f"=Konten!$A$2:$A${len(accounts) + 1}", allow_blank=True,
                        showErrorMessage=False)
    ws.add_data_validation(dv)
    dv.add("C2:D501")
    if vat:
        codes = [c for c, v in mwst.CODES.items()
                 if not (mwst.config(book)["methode"] == "saldo" and v.kind in ("vorsteuer", "investition"))]
        mc = wb.create_sheet("MWST-Codes")
        mc.append(["Code", "Bedeutung"])
        for c in codes:
            mc.append([c, mwst.CODES[c].label])
        mc.column_dimensions["B"].width = 60
        dvm = DataValidation(type="list", formula1=f"='MWST-Codes'!$A$2:$A${len(codes) + 1}", allow_blank=True)
        ws.add_data_validation(dvm)
        dvm.add("F2:F501")

    guide = wb.create_sheet("Anleitung")
    lines = [f"Buchungen für {book.settings.firma}", "",
             "Jede Zeile ist eine Buchung: Datum, Beschreibung, Soll (wohin), Haben (woher), Betrag.",
             "Konten im Dropdown wählen oder nur die Nummer eintragen (z.B. 6500).",
             "Leeres Datum = wie die Zeile darüber. Beträge positiv, mit Punkt oder Komma.",
             ("MWST: Code wählen, dann ist der Betrag brutto (inkl. MWST)." if vat else ""),
             "Die Beispielzeile löschen.", "",
             "Importieren: aeradex → Journal → «Excel importieren». Die Zeilen erscheinen zuerst im Raster;",
             "gebucht wird erst mit «Alle buchen» — alles oder nichts."]
    for line in lines:
        guide.append([line])
    guide["A1"].font = Font(bold=True, size=13)
    guide.column_dimensions["A"].width = 100
    thin = Side(style="thin", color="D6D3D1")
    for cell in ws[1]:
        cell.border = Border(bottom=thin)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
